from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE_TYPE,MSO_SHAPE
ROOT=Path(__file__).parent
exec((ROOT/'build_journal.py').read_text(encoding='utf8').split('raw=json.loads')[0])
OUT=OUT.parent/'单片机课设_学术完整答辩_实拍配图版'
prs=Presentation(OUT/'超声波测距与分级报警_实拍配图版.pptx')
def call(s,p,crop,point,label,lx,ly,w=1.8):
    l,t,r,b=crop
    px=p.left/914400+(point[0]-l)/(r-l)*p.width/914400
    py=p.top/914400+(point[1]-t)/(b-t)*p.height/914400
    # Native editable leader, endpoint and label; source photo stays intact.
    start=(lx+w/2,ly+.17)
    dk.connector(s,start,(px,py),color=C['white'],width=3.3,arrow=False)
    dk.connector(s,start,(px,py),color=C['green'],width=1.7,arrow=True)
    dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(px-.04),Inches(py-.04),Inches(.08),Inches(.08))
    dot.fill.solid();dot.fill.fore_color.rgb=C['white'];dot.line.color.rgb=C['green']
    box(s,lx,ly,w,.36,'green')
    txt(s,lx+.07,ly+.015,w-.1,.32,label,13,True,'white')
plans={
3:((0,.18,1,.87),[
((.61,.54),'STC单片机',3.75,2.48,1.65),
((.17,.60),'LCD1602显示屏',1.12,3.05,1.9),
((.55,.70),'SR04超声波模块',2.95,5.83,2),
((.738,.47),'蜂鸣器',5.7,2.55,1.4),
((.925,.715),'K3 / K4独立按键',5.55,5.9,2.05)]),
4:((.25,.20,.99,.94),[
((.55,.40),'P1端口区',1.04,2.2,1.55),
((.56,.62),'P3端口区',1.04,4.05,1.55),
((.78,.62),'STC89C516RD+',3.25,5.95,1.9)]),
17:((.27,.70,.78,.96),[
((.615,.786),'K3：阈值 +5cm',3.35,2.1,2.05),
((.705,.786),'K4：阈值 −5cm',5.65,2.1,2.05)]),
21:((.08,.29,.54,.52),[
((.267,.371),'J7控制端',1.25,2.17,1.65),
((.215,.452),'无源蜂鸣器',4.95,5.75,1.8)]),
29:((0,0,1,1),[
((.733,.205),'LCD：距离 / 阈值',5.35,2.35,2.15),
((.278,.85),'SR04超声波模块',1.2,5.83,2.1)]),
31:((.06,.48,.99,.93),[
((.287,.64),'距离37cm / 阈值30cm',1.15,2.18,2.65),
((.535,.859),'手持SR04模块',4.9,5.85,2)]),
}
for pg,(crop,labels) in plans.items():
    s=prs.slides[pg-1];p=next(x for x in s.shapes if x.shape_type==MSO_SHAPE_TYPE.PICTURE)
    for pt,lab,x,y,w in labels:call(s,p,crop,pt,lab,x,y,w)
dest=OUT/'超声波测距与分级报警_实拍标注版.pptx';prs.save(dest);print(dest)

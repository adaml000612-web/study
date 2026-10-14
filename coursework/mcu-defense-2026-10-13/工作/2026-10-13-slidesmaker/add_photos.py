from pathlib import Path
import shutil,json
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
ROOT=Path(__file__).parent
# Reuse the deck's native typography and primitives without rebuilding the source deck.
exec((ROOT/'build_journal.py').read_text(encoding='utf8').split('raw=json.loads')[0])
BASE=OUT
OUT=BASE.parent/'单片机课设_学术完整答辩_实拍配图版'
OUT.mkdir(exist_ok=True)
assets=OUT/'原始照片';assets.mkdir(exist_ok=True)
origin=Path(r'D:\xwechat_files\wxid_s1uv33cr8u5312_a321\temp\RWTemp\2026-10\d4fa906ec668d8d90cf367aa992b3dbe')
names=['abfb41d165721084cec43e87f0e35e86','9d276ee9919f677eb68c59e2b062cc36','6894c3f48a0721fbb92a7a92e291cf03','fcd40e520d316ab6c36e94f9cb47e4a7','734ff403d0d6c046576d0aa6daef831b']
for j,n in enumerate(names,1):shutil.copy2(origin/(n+'.jpg'),assets/f'{j:02}.jpg')
prs=Presentation(BASE/'超声波测距与分级报警_完整学术汇报.pptx')
def picture(s,num,x,y,w,h,crop=(0,0,1,1)):
    p=assets/f'{num:02}.jpg'
    with Image.open(p) as im:iw,ih=im.size
    l,t,r,b=crop;ratio=iw*(r-l)/(ih*(b-t));dw=min(w,h*ratio);dh=dw/ratio
    sh=s.shapes.add_picture(str(p), Inches(x+(w-dw)/2), Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))
    sh.crop_left=l;sh.crop_top=t;sh.crop_right=1-r;sh.crop_bottom=1-b
    sh._element.nvPicPr.cNvPr.set('descr',f'用户提供的实验实拍照片{num}；仅在PPT内裁切展示，原始文件保留。')
    return sh
spec={
3:(1,(0,.18,1,.87),'整套装置：开发板、SR04、LCD与连接线',[
('测量与运算','SR04输出回波信号，STC89C516RD+完成计时、厘米换算和报警状态判断。'),
('显示与交互','外接LCD显示测量信息；板载独立按键与蜂鸣模块分别承担阈值输入和声音提示。')]),
4:(2,(.25,.20,.99,.94),'单片机与端口接线近照；接口对应以源码为准',[
('超声波与LCD','Trig接P1.0，Echo接P1.1；LCD数据总线接P0，RS、RW、E分别接P2.6、P2.5、P2.7。'),
('按键与蜂鸣器','K3、K4接P3.2、P3.3；P1.5连接蜂鸣模块J7。照片展示端口区域，遮挡处结合接线表说明。')]),
17:(3,(.27,.70,.78,.96),'独立按键区：上排黑色按键K1～K4，使用K3/K4',[
('阈值输入','K3增加5cm，K4减少5cm；有效范围5～200cm，初始阈值30cm。使用黑色独立按键。'),
('消抖与保存','稳定3次扫描后执行一次；松手才可再次触发。阈值存于RAM，复位恢复30cm，LCD第二行显示Limit。')]),
21:(3,(.08,.29,.54,.52),'蜂鸣器及J7控制连接局部（由原照片裁切）',[
('实物驱动路径','P1.5连接J7控制端，经PNP驱动无源蜂鸣器；高电平关闭，低电平导通。'),
('程序产生音调','T2装载0xFF1A，约每250μs翻转输出一次，完整周期约500μs，名义音调约2kHz。')]),
29:(4,(0,0,1,1),'升级后运行实拍：LCD同时显示Dist与Limit',[
('完整装置运行','照片可见超声波模块、开发板和已点亮的LCD；界面同时显示实时距离与Limit 030cm。'),
('版本与验证记录','基础整合、阈值、回差、分级版代码规模依次1178、1824、2016、2286，构建均0错0警告；现有反馈支持基本功能试用。')]),
31:(5,(.06,.48,.99,.93),'操作现场：手持超声波模块，LCD显示距离37cm、阈值30cm',[
('现场演示顺序','改变目标距离→按键调整阈值→观察报警进入与退出→比较三档鸣叫节奏→展示无回波状态。'),
('照片与演示说明','实拍展示升级后的距离与阈值界面。声音节奏由现场演示说明；精度需尺距对照，尚未加入滤波、温补或EEPROM保存。')])}
for page,(num,crop,caption,items) in spec.items():
    s=prs.slides[page-1]
    for sh in list(s.shapes):
        if sh.top>=Inches(1.95) and sh.top<Inches(6.95):s.shapes._spTree.remove(sh._element)
    portrait=page==4
    if portrait:
        picture(s,num,1,2.05,4.0,4.5,crop);bullets(s,items,x=5.35,y=2.12,w=6.9,h=4.45)
    else:
        picture(s,num,.97,2.08,6.8,4.25,crop);bullets(s,items,x=8.05,y=2.12,w=4.25,h=4.65)
    txt(s,1,6.58,11.25,.36,caption,13,False,'mute')
    old=s.notes_slide.notes_text_frame.text
    dk.speaker_notes(s,old+'\n\n本页配图：'+caption+'。用户本轮提供照片，不能据此认定早期独立测试、精度或声音节奏已获图像验证。')
dest=OUT/'超声波测距与分级报警_实拍配图版.pptx';prs.save(dest)
(OUT/'逐页讲稿.md').write_text('\n\n'.join(f'## {i+1}. {s.shapes.title.text if s.shapes.title else ""}\n\n{s.notes_slide.notes_text_frame.text}' for i,s in enumerate(prs.slides)),encoding='utf8')
(OUT/'配图记录.json').write_text(json.dumps({'pages':{str(p):{'photo':v[0],'caption':v[2]} for p,v in spec.items()},'originals_preserved':True,'previous_deck_unchanged':True},ensure_ascii=False,indent=2),encoding='utf8')
print(dest)
